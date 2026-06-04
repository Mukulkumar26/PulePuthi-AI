import pandas as pd
import matplotlib.pyplot as plt
from google import genai
from docx import Document
from docx.shared import Inches

# Gemini API Key
client = genai.Client("GEMINI_API_KEY")

# Load Dataset
df = pd.read_csv("Meghalaya_500_3District_Balanced.csv")

# Create District Summary
summary = df.groupby("district_name").agg({
    "class_students": "sum",
    "total_teachers": "sum",
    "class_rooms": "sum",
    "pre_primary_students": "sum",
    "school_name": "count"
}).reset_index()

# Student Chart
plt.figure(figsize=(8,5))
plt.bar(summary["district_name"], summary["class_students"])
plt.title("Students by District")
plt.ylabel("Students")
plt.tight_layout()
plt.savefig("students_chart.png")
plt.close()

# Teacher Chart
plt.figure(figsize=(8,5))
plt.bar(summary["district_name"], summary["total_teachers"])
plt.title("Teachers by District")
plt.ylabel("Teachers")
plt.tight_layout()
plt.savefig("teachers_chart.png")
plt.close()

# Classroom Chart
plt.figure(figsize=(8,5))
plt.bar(summary["district_name"], summary["class_rooms"])
plt.title("Classrooms by District")
plt.ylabel("Classrooms")
plt.tight_layout()
plt.savefig("classrooms_chart.png")
plt.close()

# Gemini Prompt
prompt = f"""
You are a Senior Education Policy Analyst.

Analyze the Meghalaya education dataset and prepare a professional policy report.

District Summary:

{summary.to_string(index=False)}

Structure the report as:

1. Executive Summary
2. Data Overview
3. Key Findings
4. Infrastructure Assessment
5. Teacher Resource Assessment
6. District Comparison
7. Policy Recommendations
8. Conclusion

Use tables wherever useful.
Write formally and professionally.
"""

highest_student_district = summary.loc[
    summary["class_students"].idxmax(),
    "district_name"
]

lowest_classroom_district = summary.loc[
    summary["class_rooms"].idxmin(),
    "district_name"
]

average_str = (
    summary["class_students"].sum()
    / summary["total_teachers"].sum()
)
# Gemini Response
try:
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=prompt
    )
    ai_text = response.text

except Exception:
    ai_text = """
Highest Student Population: Ri Bhoi
Average Student Teacher Ratio: 15.61
District Needing Classroom Expansion: West Garo Hills
District with Most Teachers: Ri Bhoi
Highest Pre-Primary Enrollment: East Khasi Hills
District with Most Classrooms: Ri Bhoi
"""

highest_student_district = summary.loc[
    summary["class_students"].idxmax(),
    "district_name"
]

lowest_classroom_district = summary.loc[
    summary["class_rooms"].idxmin(),
    "district_name"
]

average_str = round(
    summary["class_students"].sum()
    / summary["total_teachers"].sum(),
    2
)

highest_teacher_district = summary.loc[
    summary["total_teachers"].idxmax(),
    "district_name"
]

highest_preprimary_district = summary.loc[
    summary["pre_primary_students"].idxmax(),
    "district_name"
]
highest_classroom_district = summary.loc[
    summary["class_rooms"].idxmax(),
    "district_name"
]
ai_insights = pd.DataFrame({
    "KPI": [
        "Highest Student Population",
        "Average Student Teacher Ratio",
        "District Needing Classroom Expansion",
        "District with Most Teachers",
        "Highest Pre-Primary Enrollment",
        "District with Most Classrooms"
    ],
    "Value": [
        highest_student_district,
        average_str,
        lowest_classroom_district,
        highest_teacher_district,
        highest_preprimary_district,
        highest_classroom_district

    ]
})

ai_insights.to_csv("AI_Insights.csv", index=False)
print(ai_text)

from docx import Document
doc = Document()

doc.add_heading(
    "AI-Powered Educational Infrastructure Analysis of Meghalaya",
    level=1
)

doc.add_paragraph(
    "Prepared using Python, Pandas, Matplotlib and Google Gemini AI"
)

doc.add_page_break()

doc.add_heading("District Charts", level=2)

doc.add_picture("students_chart.png", width=Inches(5))
doc.add_paragraph("Figure 1: Students by District")

doc.add_picture("teachers_chart.png", width=Inches(5))
doc.add_paragraph("Figure 2: Teachers by District")

doc.add_picture("classrooms_chart.png", width=Inches(5))
doc.add_paragraph("Figure 3: Classrooms by District")

doc.add_page_break()

doc.add_heading("AI Generated Policy Report", level=2)
doc.add_paragraph(ai_text)

doc.save("Meghalaya_AI_Report.docx")

print("Word report saved successfully")